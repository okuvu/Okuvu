#!/usr/bin/env python3
"""Markdown -> docx for the AWS migration plan. Handles headings, tables,
bullets, bold, blockquotes, horizontal rules. No dependencies beyond python-docx."""
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC, DST = sys.argv[1], sys.argv[2]

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(10.5)
for a, b in (('Heading 1', 20), ('Heading 2', 14), ('Heading 3', 11.5)):
    doc.styles[a].font.name = 'Calibri'
    doc.styles[a].font.size = Pt(b)
    doc.styles[a].font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)

def add_runs(par, text, bold_all=False):
    # split on **bold** and `code`
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith('`') and part.endswith('`'):
            r = par.add_run(part[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(9.5)
        else:
            r = par.add_run(part)
        if bold_all:
            r.bold = True
    return par

lines = open(SRC, encoding='utf-8').read().splitlines()
i = 0
first_heading_done = False
while i < len(lines):
    ln = lines[i]
    s = ln.strip()

    if not s:
        i += 1; continue

    if s == '---':
        i += 1; continue

    # table
    if s.startswith('|') and i + 1 < len(lines) and set(lines[i+1].strip().replace('|','').replace(':','')) <= {'-'}:
        header = [c.strip() for c in s.strip('|').split('|')]
        i += 2
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
            i += 1
        t = doc.add_table(rows=1, cols=len(header))
        t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = t.rows[0].cells
        for j, h in enumerate(header):
            hdr[j].text = ''
            p = hdr[j].paragraphs[0]
            add_runs(p, h, bold_all=True)
            p.runs and setattr(p.runs[0].font, 'size', Pt(9.5))
            shade(hdr[j], 'EDEDED')
        for row in rows:
            cells = t.add_row().cells
            for j, c in enumerate(row[:len(header)]):
                cells[j].text = ''
                p = cells[j].paragraphs[0]
                add_runs(p, c)
                for r in p.runs:
                    r.font.size = Pt(9.5)
        doc.add_paragraph()
        continue

    m = re.match(r'^(#{1,4})\s+(.*)$', s)
    if m:
        level = len(m.group(1))
        text = m.group(2)
        if level == 1 and not first_heading_done:
            h = doc.add_heading(level=0)
            add_runs(h, text)
            first_heading_done = True
        else:
            h = doc.add_heading(level=min(level, 3))
            add_runs(h, text)
        i += 1; continue

    if s.startswith('> '):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        add_runs(p, s[2:])
        for r in p.runs:
            r.italic = True
            r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        i += 1; continue

    mb = re.match(r'^[-*]\s+(.*)$', s)
    if mb:
        p = doc.add_paragraph(style='List Bullet')
        add_runs(p, mb.group(1))
        i += 1; continue

    mn = re.match(r'^(\d+)\.\s+(.*)$', s)
    if mn:
        p = doc.add_paragraph(style='List Number')
        add_runs(p, mn.group(2))
        i += 1; continue

    p = doc.add_paragraph()
    add_runs(p, s)
    i += 1

doc.save(DST)
print("wrote", DST)
